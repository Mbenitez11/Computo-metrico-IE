"""Independent exports: canonical JSON, explicit formulas, local HTML and findings."""
import html
import json
import shutil
from collections import defaultdict
from pathlib import Path

from .model import save

SHEETS = ['Resumen general', 'Computo consolidado', 'Por sistema', 'Por sector y nivel',
          'Mediciones originales', 'Materiales derivados', 'Alimentadores y conductores',
          'Observaciones', 'Reglas y supuestos', 'Trazabilidad', 'Indice de planos']


def text(value):
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False)
    # Formula injection from source text must not become workbook code.
    if isinstance(value, str) and value[:1] in ('=', '+', '-', '@'):
        return "'" + value
    return value


def write_excel(data, path):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties
    workbook = Workbook(); workbook.remove(workbook.active)
    workbook.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    sheets = {name: workbook.create_sheet(name) for name in SHEETS}
    def table(name, headers, rows):
        sheet = sheets[name]
        sheet.append(headers)
        for row in rows:
            sheet.append([text(v) for v in row])
        sheet.freeze_panes = 'B2'; sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(name='Aptos', bold=True, color='FFFFFF', size=11)
            cell.fill = PatternFill('solid', fgColor='173C50')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[1].height = 32
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Aptos', size=10)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='F0F5F7')
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.000'
        for col in range(1, sheet.max_column + 1):
            values = [len(str(sheet.cell(r, col).value or '')) for r in range(1, min(sheet.max_row, 80) + 1)]
            sheet.column_dimensions[get_column_letter(col)].width = min(52, max(15, max(values, default=0) * .75))
        for row in sheet.iter_rows(min_row=2):
            lines = max((max(1, (len(str(c.value or '')) + int(sheet.column_dimensions[c.column_letter].width) - 1)
                             // max(1, int(sheet.column_dimensions[c.column_letter].width))) for c in row), default=1)
            sheet.row_dimensions[row[0].row].height = min(180, max(21, 15 * lines + 6))
            for cell in row:
                if isinstance(cell.value, (int, float)) and cell.value == int(cell.value):
                    cell.number_format = '#,##0'
        sheet.print_options.horizontalCentered = True
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        # Wide audit tables paginate horizontally instead of shrinking to unreadable text.
        sheet.page_setup.fitToWidth = 0 if sheet.max_column > 11 else 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = '1:1'
    qs = data.get('quantities', [])
    trace_fields = ['id', 'project', 'system', 'sector', 'level', 'circuit', 'material_key', 'specification',
                    'description', 'unit', 'quantity', 'method', 'review_status', 'drawing', 'revision',
                    'source_refs', 'formula', 'rule_id', 'evidence', 'confidence', 'notes', 'calibration_id', 'basis_id',
                    'provision_by', 'installation_by']
    table('Trazabilidad', trace_fields, [[q.get(k) for k in trace_fields] for q in qs])
    trace_rows = {q['id']: i + 2 for i, q in enumerate(qs)}
    table('Mediciones originales', ['ID', 'Sistema', 'Sector', 'Nivel', 'Descripcion', 'Unidad', 'Cantidad', 'Metodo', 'Estado', 'Fuentes'],
          [[q['id'], q['system'], q['sector'], q['level'], q['description'], q['unit'], q['quantity'], q['method'], q['review_status'], q['source_refs']]
           for q in qs if not q['basis_id']])
    derived = [q for q in qs if q['basis_id']]
    table('Materiales derivados', ['ID', 'Base', 'Regla', 'Cantidad', 'Unidad', 'Estado', 'Formula', 'Evidencia'],
          [[q['id'], q['basis_id'], q['rule_id'], q['quantity'], q['unit'], q['review_status'], q['formula'], q['evidence']] for q in derived])
    rules = {r['id']: r for r in data.get('rules', [])}
    rule_rows = []; variable_rows = {}
    for rule in data.get('rules', []):
        for key, value in rule.get('variables', {}).items():
            rule_rows.append([rule['id'], rule.get('version'), rule.get('kind'), key, value,
                              rule.get('variable_units', {}).get(key), rule.get('status'), rule.get('source'),
                              rule.get('evidence'), rule.get('formula'), rule.get('conditions')])
            variable_rows[(rule['id'], key)] = len(rule_rows) + 1
    table('Reglas y supuestos', ['ID regla', 'Version', 'Tipo', 'Variable', 'Valor', 'Unidad', 'Aprobacion', 'Fuente', 'Evidencia', 'Formula declarada', 'Condiciones'], rule_rows)
    # Formula ownership: rules + original trace -> derived -> derived trace -> summary.
    for row, q in enumerate(derived, 2):
        if q['quantity'] is None or q['basis_id'] not in trace_rows:
            continue
        rule = rules[q['rule_id']]; base = f"'Trazabilidad'!K{trace_rows[q['basis_id']]}"
        def v(key):
            return f"'Reglas y supuestos'!E{variable_rows[(rule['id'], key)]}"
        if rule['kind'] == 'conductor':
            formula = f"=({base}+{v('vertical')}+{v('reserve')})*{v('count')}*(1+{v('waste')})"
        elif rule['kind'] == 'supports':
            formula = f"=ROUNDUP({base}/{v('spacing')},0)+{v('endpoints')}"
        else:
            formula = f"={base}*{v('fraction')}"
        sheets['Materiales derivados'].cell(row, 4, formula)
        sheets['Trazabilidad'].cell(trace_rows[q['id']], 11, f"='Materiales derivados'!D{row}")
    cons = data.get('consolidated', [])
    table('Computo consolidado', ['ID', 'Sistema', 'Material', 'Especificacion', 'Descripcion', 'Unidad', 'Metodo', 'Cantidad confirmada', 'Partidas originales'],
          [[c['id'], c['system'], c['material_key'], c['specification'], c['description'], c['unit'], c['method'], c['quantity'], c['source_ids']] for c in cons])
    for row, c in enumerate(cons, 2):
        terms = [f"'Trazabilidad'!K{trace_rows[qid]}" for qid in c['source_ids']]
        sheets['Computo consolidado'].cell(row, 8, '=' + '+'.join(terms))
    table('Por sistema', ['Sistema', 'Material', 'Unidad', 'Metodo', 'Cantidad confirmada', 'ID consolidado'],
          [[c['system'], c['material_key'], c['unit'], c['method'], c['quantity'], c['id']] for c in cons])
    for row in range(2, len(cons) + 2):
        sheets['Por sistema'].cell(row, 5, f"='Computo consolidado'!H{row}")
    sector_groups = defaultdict(list)
    for q in qs:
        if q['review_status'] == 'confirmed':
            sector_groups[(q['sector'], q['level'], q['system'], q['material_key'], q['unit'], q['method'])].append(q)
    table('Por sector y nivel', ['Sector', 'Nivel', 'Sistema', 'Material', 'Unidad', 'Metodo', 'Cantidad confirmada', 'IDs'],
          [[*key, sum(q['quantity'] for q in group), [q['id'] for q in group]] for key, group in sorted(sector_groups.items())])
    for row, (_, group) in enumerate(sorted(sector_groups.items()), 2):
        sheets['Por sector y nivel'].cell(row, 7, '=' + '+'.join(f"'Trazabilidad'!K{trace_rows[q['id']]}" for q in group))
    conductors = [q for q in derived if rules[q['rule_id']]['kind'] == 'conductor']
    table('Alimentadores y conductores', ['ID', 'Circuito', 'Material', 'Especificacion', 'Unidad', 'Cantidad', 'Estado', 'Base', 'Regla'],
          [[q['id'], q['circuit'], q['material_key'], q['specification'], q['unit'], q['quantity'], q['review_status'], q['basis_id'], q['rule_id']] for q in conductors])
    for row, q in enumerate(conductors, 2):
        if q['quantity'] is not None:
            sheets['Alimentadores y conductores'].cell(row, 6, f"='Trazabilidad'!K{trace_rows[q['id']]}")
    obs_fields = ['id', 'severity', 'code', 'documents', 'items', 'values', 'detail', 'possible_cause', 'impact', 'action']
    table('Observaciones', obs_fields, [[o.get(k) for k in obs_fields] for o in data.get('observations', [])])
    table('Indice de planos', ['ID', 'Archivo', 'SHA256', 'Formato', 'Especialidad indicativa', 'Revision', 'Estado', 'Firma DWG', 'Hojas'],
          [[d['id'], d['name'], d['sha256'], d['extension'], d['system_hint'], d.get('revision'), d['status'], d.get('dwg_signature'), d.get('sheets')]
           for d in data['documents']])
    table('Resumen general', ['Concepto', 'Valor', 'Alcance'],
          [['Proyecto', data['project'], 'Cómputo métrico sin precios'],
           ['Documentos', len(data['documents']), 'Ver índice y estados'],
           ['Partidas confirmadas', sum(q['review_status'] == 'confirmed' for q in qs), 'Solo medidas o calculadas revisadas'],
           ['Partidas sin confirmar', sum(q['review_status'] != 'confirmed' for q in qs), 'Incluye pendientes, estimadas y conflictos'],
           ['Observaciones bloqueantes', sum(o['severity'] == 'blocking' for o in data.get('observations', [])), 'Ver Observaciones'],
           ['Cobertura', 'Sin conteo verificado' if not qs else 'Parcial: revisar alcance seleccionado', 'No sumar cantidades de unidades distintas'],
           ['Actualización', 'Recalcular desde config.json', 'El visor y JSON reflejan la última ejecución; Excel no cambia esos archivos']]
          + [[c['system'], c['status'], c.get('evidence') or 'Revisar documentación y elementos necesarios, incluidos los provistos por terceros'] for c in data.get('coverage', [])])
    # Input variables are editable; provenance and computed cells have distinct fonts.
    for row in sheets['Reglas y supuestos'].iter_rows(min_row=2):
        row[4].font = Font(name='Aptos', color='0055AA', size=10)
        row[4].fill = PatternFill('solid', fgColor='FFF3CB')
    workbook.save(path)
    # Reopen to validate package, sheet names and broken-reference text, without evaluating formulas.
    check = load_workbook(path, data_only=False)
    if check.sheetnames != SHEETS:
        raise ValueError('Exportación perdió hojas')
    for sheet in check:
        for row in sheet:
            if any(c.data_type == 'f' and '#REF!' in c.value for c in row):
                raise ValueError('Fórmula con referencia rota')
    check.close()


def write_report(data, path):
    docs = {d['id']: d['name'] for d in data['documents']}
    qs = data.get('quantities', [])
    lines = [f"# Revisión — {data['project']}", '',
             f"Documentos: {len(docs)}. Partidas: {len(qs)}. Confirmadas: {sum(q['review_status']=='confirmed' for q in qs)}.", '',
             'Las cantidades del índice documental no son conteos verificados. Los elementos pendientes conservan cantidad vacía. No se suman unidades distintas.', '']
    for o in data.get('observations', []):
        lines += [f"## {o['severity']} · {o['code']}", '', o['detail'], '',
                  'Documentos: ' + ', '.join(docs.get(i, i) for i in o['documents']),
                  'Valores: ' + json.dumps(o.get('values', []), ensure_ascii=False),
                  'Causa posible: ' + o['possible_cause'], 'Impacto: ' + o['impact'],
                  'Acción: ' + o['action'], '']
    Path(path).write_text('\n'.join(lines), encoding='utf-8')


def write_viewer(data, path):
    template = Path(__file__).resolve().parents[2] / 'assets' / 'viewer.html'
    payload = json.dumps(data, ensure_ascii=False, allow_nan=False).replace('<', '\\u003c').replace('>', '\\u003e').replace('&', '\\u0026')
    Path(path).write_text(template.read_text(encoding='utf-8').replace('__EQT_DATA__', payload), encoding='utf-8')


def export_all(data, out, *, index_directory=None, make_excel=True, filename='quantities.json'):
    out = Path(out).resolve(); out.mkdir(parents=True, exist_ok=True)
    if index_directory:
        source = Path(index_directory).resolve() / 'images'
        destination = out / 'images'
        if source.exists() and source != destination:
            shutil.copytree(source, destination, dirs_exist_ok=True)
    save(out / filename, data)
    write_report(data, out / 'observations.md')
    write_viewer(data, out / 'viewer.html')
    if make_excel:
        write_excel(data, out / 'takeoff.xlsx')
