"""Read-only document adapters; extraction never implies electrical identity."""
import csv
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from .model import SCHEMA, SYSTEMS, uid, sha256, issue


def classify(name):
    text = name.casefold()
    patterns = [('unifilares', 'unifilar'), ('iluminacion', 'ilum'), ('tomacorrientes', 'toma'),
                ('bandejas', 'bandeja'), ('puesta_tierra', 'tierra'), ('fuerza', 'bomb'),
                ('fuerza', 'termo'), ('aire_acondicionado', ' - aa '), ('cuadro_cargas', 'cargas'),
                ('referencia_computo', 'computo')]
    return next((system for system, word in patterns if word in text), 'sin_clasificar')


def pdf_extract(path, doc, out):
    import fitz
    pages, refs, findings = [], [], []
    with fitz.open(path) as pdf:
        for n, page in enumerate(pdf):
            original_rotation = page.rotation
            page.set_rotation(0)  # in-memory only; same frame for text, paths and background
            pid = uid(doc['id'], 'page', n + 1)
            words = page.get_text('words', sort=True)
            drawings = page.get_drawings()
            text = page.get_text()
            kind = 'vector' if drawings else ('text' if words else 'raster_or_empty')
            image_name = f'images/{pid}.png'
            dest = out / image_name; dest.parent.mkdir(parents=True, exist_ok=True)
            page.get_pixmap(matrix=fitz.Matrix(1.25, 1.25), alpha=False).save(str(dest))
            pages.append({'id': pid, 'document_id': doc['id'], 'page': n + 1,
                          'width': page.rect.width, 'height': page.rect.height,
                          'original_rotation': original_rotation, 'coordinate_frame': 'unrotated_pdf_points',
                          'image': image_name, 'kind': kind, 'text': text,
                          'scale_notes': re.findall(r'1\s*:\s*\d+', text), 'path_count': len(drawings)})
            for j, w in enumerate(words):
                refs.append({'id': uid(pid, 'word', j), 'document_id': doc['id'], 'page_id': pid,
                             'kind': 'word', 'bbox': list(w[:4]), 'text': w[4],
                             'locator': f'page={n+1};word={j}'})
            for j, d in enumerate(drawings):
                commands = []
                for item in d['items']:
                    if item[0] == 'l':
                        commands.append(['l', list(item[1]), list(item[2])])
                    elif item[0] == 'c':
                        commands.append(['c'] + [list(p) for p in item[1:]])
                    elif item[0] == 're':
                        commands.append(['re', list(item[1])])
                    elif item[0] == 'qu':
                        q = item[1]
                        commands.append(['qu', list(q.ul), list(q.ur), list(q.lr), list(q.ll)])
                    else:
                        commands.append(['unsupported', str(item)])
                # closePath can add an unlisted segment. Include only a single continuous line chain.
                if d.get('closePath') and commands and all(c[0] == 'l' for c in commands):
                    if all(commands[k][2] == commands[k+1][1] for k in range(len(commands)-1)):
                        if commands[-1][2] != commands[0][1]:
                            commands.append(['l', commands[-1][2], commands[0][1]])
                    else:
                        commands.append(['unsupported', 'multiple closed subpaths'])
                refs.append({'id': uid(pid, 'path', j), 'document_id': doc['id'], 'page_id': pid,
                             'kind': 'path', 'bbox': list(d['rect']), 'commands': commands,
                             'layer': d.get('layer'), 'color': d.get('color'),
                             'locator': f'page={n+1};path={j}'})
            if kind == 'raster_or_empty' or page.get_images():
                findings.append(issue('RASTER_REVIEW', 'warning', [doc['id']],
                                      f'Página {n+1}: contenido raster o mixto; vectores no prueban cobertura total.',
                                      'Revisar imagen; OCR y símbolos raster requieren adaptación validada.'))
            findings.append(issue('SCALE_UNVERIFIED', 'blocking', [doc['id']],
                                  f'Página {n+1}: escala geométrica todavía no aprobada.',
                                  'Calibrar por cota conocida y región antes de medir.'))
    return pages, refs, findings


def excel_extract(path, doc):
    cells, sheets, findings = [], [], []
    if path.suffix.lower() == '.xls':
        import xlrd
        book = xlrd.open_workbook(path)
        for sheet in book.sheets():
            sheets.append({'name': sheet.name, 'rows': sheet.nrows, 'columns': sheet.ncols})
            for row in range(sheet.nrows):
                for col in range(sheet.ncols):
                    cell = sheet.cell(row, col)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_ERROR:
                        findings.append(issue('SOURCE_CELL_ERROR', 'warning', [doc['id']],
                                              f'{sheet.name} fila {row+1}, columna {col+1}: error Excel {value}',
                                              'Revisar fórmula en aplicación de origen.'))
                        value = None
                    cells.append(cell_record(doc, sheet.name, row + 1, col + 1, value,
                                             storage='xls_cached_value', cell_type=cell.ctype))
    else:
        from openpyxl import load_workbook
        values = load_workbook(path, read_only=True, data_only=True)
        formulas = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in formulas:
                sheets.append({'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column})
                vs = values[sheet.title]
                for row in sheet:
                    for cell in row:
                        if cell.value is None:
                            continue
                        formula = cell.value if cell.data_type == 'f' else None
                        value = vs.cell(cell.row, cell.column).value if formula else cell.value
                        if value is not None and not isinstance(value, (str, int, float, bool)):
                            value = value.isoformat() if hasattr(value, 'isoformat') else str(value)
                        cells.append(cell_record(doc, sheet.title, cell.row, cell.column, value, formula=formula))
                        if formula and value is None:
                            findings.append(issue('FORMULA_NO_CACHE', 'warning', [doc['id']],
                                                  f'{sheet.title}!{cell.coordinate}: fórmula sin caché.',
                                                  'Recalcular copia con Excel/LibreOffice; no reemplazar por cero.'))
        finally:
            values.close(); formulas.close()
    return cells, sheets, findings


def cell_record(doc, sheet, row, col, value, **extra):
    from openpyxl.utils import get_column_letter
    address = f'{get_column_letter(col)}{row}'
    return {'id': uid(doc['id'], sheet, address), 'document_id': doc['id'], 'sheet': sheet,
            'address': address, 'row': row, 'column': col, 'value': value, **extra}


def index_folder(root, out, project):
    root = Path(root).resolve(); out = Path(out).resolve()
    if not root.exists():
        raise ValueError('No existe la entrada')
    if out == root or (root.is_dir() and out.is_relative_to(root)):
        raise ValueError('La salida debe estar fuera de la carpeta de entrada')
    files = [root] if root.is_file() else sorted(p for p in root.rglob('*') if p.is_file())
    out.mkdir(parents=True, exist_ok=True)
    result = {'schema': SCHEMA, 'project': project, 'documents': [], 'pages': [], 'references': [],
              'cells': [], 'observations': [], 'quantities': [], 'rules': [], 'consolidated': [],
              'coverage': [{'system': s, 'status': 'not_reviewed', 'evidence': ''} for s in SYSTEMS]}
    hashes = {}
    for path in files:
        digest = sha256(path)
        doc = {'id': uid('document', digest), 'sha256': digest, 'name': path.name,
               'path': str(path), 'extension': path.suffix.lower(), 'bytes': path.stat().st_size,
               'system_hint': classify(path.name), 'revision_hint': (re.findall(r'RV\d+', path.name, re.I) or [None])[0],
               'revision': None, 'status': 'indexed', 'sheets': []}
        if digest in hashes:
            result['observations'].append(issue('DUPLICATE_FILE', 'warning', [doc['id']],
                                                f'Archivo repetido por hash: {path.name}; {hashes[digest]}',
                                                'Elegir una copia autoritativa.'))
            continue
        hashes[digest] = str(path); result['documents'].append(doc)
        try:
            ext = path.suffix.lower()
            if ext == '.pdf':
                pages, refs, obs = pdf_extract(path, doc, out)
                result['pages'].extend(pages); result['references'].extend(refs); result['observations'].extend(obs)
            elif ext in ('.xls', '.xlsx', '.xlsm'):
                cells, sheets, obs = excel_extract(path, doc)
                result['cells'].extend(cells); doc['sheets'] = sheets; result['observations'].extend(obs)
            elif ext in ('.txt', '.md'):
                doc['text'] = path.read_text(encoding='utf-8-sig', errors='replace')
            elif ext == '.docx':
                with zipfile.ZipFile(path) as z:
                    xml = ElementTree.fromstring(z.read('word/document.xml'))
                    doc['text'] = '\n'.join(t.text or '' for t in xml.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
            else:
                doc['status'] = 'pending_adapter'
                if ext == '.dwg':
                    with path.open('rb') as f:
                        doc['dwg_signature'] = f.read(6).decode('ascii', errors='replace')
                result['observations'].append(issue('ADAPTER_REQUIRED', 'blocking', [doc['id']],
                                                    f'{path.name}: no se extrajo geometría ({ext}).',
                                                    'Convertir copia a PDF vectorial validado; conservar origen, hash, versión y unidades.'))
        except Exception as exc:
            doc['status'] = 'extraction_failed'
            result['observations'].append(issue('EXTRACTION_FAILED', 'blocking', [doc['id']],
                                                f'{path.name}: {type(exc).__name__}: {exc}',
                                                'Revisar formato/dependencia; no asumir documento vacío.'))
    return result
