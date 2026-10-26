"""Synthetic acceptance fixtures; no quantities here represent a real project."""
import copy
import json
import math
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from eqt.model import meters, number, geometry_length, save, uid
from eqt.extract import index_folder
from eqt.engine import run, tag_candidates, rule_value
from eqt.export import export_all, SHEETS


class PipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import fitz
        cls.temp = tempfile.TemporaryDirectory(prefix='eqt-synthetic-')
        cls.root = Path(cls.temp.name); cls.inputs = cls.root / 'originals'; cls.inputs.mkdir()
        pdf = fitz.open(); page = pdf.new_page(width=300, height=300)
        page.insert_text((30, 25), 'SYNTHETIC FIXTURE - NOT REAL PROJECT DATA')
        page.draw_line((40, 50), (70, 90))  # 50 pt => 5 m at .1 m/pt
        page.draw_rect(fitz.Rect(100, 50, 110, 60))
        page.draw_bezier((150, 50), (160, 30), (170, 100), (180, 60))
        page.insert_text((100, 80), 'L1'); page.insert_text((150, 80), 'L1')
        page.insert_text((40, 220), '1:100')  # not used as calibration
        page2 = pdf.new_page(width=300, height=300); page2.draw_line((20, 20), (50, 60)); page2.set_rotation(90)
        pdf.save(cls.inputs / 'SYNTHETIC-IE-01-R01.pdf'); pdf.close()
        cls.index_dir = cls.root / 'index'
        cls.index = index_folder(cls.inputs, cls.index_dir, 'SYNTHETIC')
        cls.doc = cls.index['documents'][0]; cls.page = cls.index['pages'][0]
        cls.paths = [r for r in cls.index['references'] if r['page_id'] == cls.page['id'] and r['kind'] == 'path']
        cls.words = [r for r in cls.index['references'] if r.get('text') == 'L1']

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def config(self):
        return {'schema': 'eqt/1', 'documents': {self.doc['id']: {
            'sha256': self.doc['sha256'], 'family': 'IE-01', 'revision': 'R01',
            'status': 'approved', 'evidence': 'Synthetic reviewed revision'}},
            'calibrations': [{'id': 'cal', 'page_id': self.page['id'],
                              'source_sha256': self.doc['sha256'], 'p1': [40, 200], 'p2': [140, 200],
                              'real_distance': 10, 'unit': 'm', 'region': [0, 0, 300, 300],
                              'status': 'approved', 'evidence': 'Synthetic known dimension'}],
            'selections': [self.selection()], 'rules': [], 'comparisons': []}

    def selection(self, **overrides):
        s = {'id': 'route', 'kind': 'length', 'system': 'bandejas', 'sector': 'A', 'level': 'PB',
             'material_key': 'tray-200', 'description': 'Synthetic route', 'specification': '200 mm',
             'unit': 'm', 'source_refs': [self.paths[0]['id']], 'calibration_id': 'cal',
             'installation_route_verified': True, 'status': 'approved', 'evidence': 'Synthetic route review'}
        s.update(overrides); return s

    def rule(self, **overrides):
        r = {'id': 'cable', 'kind': 'conductor', 'basis_id': 'route', 'description': 'Synthetic conductor',
             'system': 'alimentadores', 'material_key': 'cu-35', 'specification': '35 mm2', 'unit': 'm',
             'formula': '(L+vertical+reserve)*count*(1+waste)',
             'variables': {'count': 3, 'vertical': 2, 'reserve': 1, 'waste': .1},
             'variable_units': {'count': 'u', 'vertical': 'm', 'reserve': 'm', 'waste': '1'},
             'conditions': {'system': 'bandejas'}, 'source': 'Synthetic unifilar', 'version': '1',
             'status': 'approved', 'evidence': 'Synthetic rule review'}
        r.update(overrides); return r

    def test_length_345_calibrated(self):
        result = run(self.index, self.config())
        self.assertAlmostEqual(result['quantities'][0]['quantity'], 5)
        self.assertEqual(result['quantities'][0]['review_status'], 'confirmed')

    def test_missing_scale_not_zero(self):
        c = self.config(); c['calibrations'] = []
        q = run(self.index, c)['quantities'][0]
        self.assertIsNone(q['quantity']); self.assertEqual(q['method'], 'pending')

    def test_scale_note_not_authorization(self):
        self.assertIn('1:100', self.page['scale_notes'])
        c = self.config(); c['calibrations'][0]['status'] = 'proposed'
        self.assertIsNone(run(self.index, c)['quantities'][0]['quantity'])

    def test_calibration_hash(self):
        c = self.config(); c['calibrations'][0]['source_sha256'] = 'wrong'
        self.assertIsNone(run(self.index, c)['quantities'][0]['quantity'])

    def test_outside_region(self):
        c = self.config(); c['calibrations'][0].update(region=[30, 180, 160, 230])
        self.assertIsNone(run(self.index, c)['quantities'][0]['quantity'])

    def test_zero_calibration(self):
        c = self.config(); c['calibrations'][0]['real_distance'] = 0
        self.assertIsNone(run(self.index, c)['quantities'][0]['quantity'])

    def test_bad_units(self):
        with self.assertRaises(ValueError): meters(3, 'pixels')
        self.assertAlmostEqual(meters(1, 'ft'), .3048)
        self.assertAlmostEqual(meters(1000, 'mm'), 1)

    def test_bad_numbers(self):
        for value in [True, -1, float('nan'), float('inf'), '12']:
            with self.assertRaises(ValueError): number(value)

    def test_exact_primitives(self):
        self.assertEqual(geometry_length([['l', [0, 0], [3, 4]]]), 5)
        self.assertEqual(geometry_length([['re', [0, 0, 3, 4]]]), 14)

    def test_curves_blocked(self):
        c = self.config(); c['selections'][0]['source_refs'] = [self.paths[2]['id']]
        self.assertIsNone(run(self.index, c)['quantities'][0]['quantity'])

    def test_route_must_be_verified(self):
        c = self.config(); c['selections'][0]['installation_route_verified'] = False
        self.assertIsNone(run(self.index, c)['quantities'][0]['quantity'])

    def test_source_revision_not_approved(self):
        c = self.config(); c['documents'][self.doc['id']]['status'] = 'proposed'
        self.assertEqual(run(self.index, c)['consolidated'], [])

    def test_count_is_one_physical_element(self):
        c = self.config(); c['selections'] = [self.selection(kind='count', unit='u', source_refs=[self.paths[1]['id']])]
        self.assertEqual(run(self.index, c)['quantities'][0]['quantity'], 1)

    def test_tag_is_only_candidate(self):
        self.assertEqual(len(tag_candidates(self.index, {'L1': 'lamp'})), 2)
        c = self.config(); c['selections'] = [self.selection(kind='count', unit='u', source_refs=[self.words[0]['id']])]
        self.assertIsNone(run(self.index, c)['quantities'][0]['quantity'])
        c['selections'][0]['tag_is_element_verified'] = True
        self.assertEqual(run(self.index, c)['quantities'][0]['quantity'], 1)

    def test_duplicate_reference_blocked(self):
        c = self.config(); c['selections'].append(self.selection(id='second'))
        r = run(self.index, c)
        self.assertEqual(r['consolidated'], [])
        self.assertTrue(all(q['review_status'] == 'duplicate_review' for q in r['quantities']))

    def test_duplicate_resolution_preserves_sources(self):
        c = self.config(); c['selections'].append(self.selection(id='second'))
        c['identity_resolutions'] = [{'item_ids': ['route', 'second'], 'keep_id': 'route', 'status': 'approved', 'evidence': 'Same synthetic route'}]
        r = run(self.index, c)
        self.assertEqual(r['consolidated'][0]['quantity'], 5)
        self.assertEqual(len(r['quantities']), 2)

    def test_estimate_separate(self):
        c = self.config(); c['selections'][0].update(kind='estimate', value=10, justification='Synthetic assumption')
        r = run(self.index, c)
        self.assertEqual(r['quantities'][0]['method'], 'estimated'); self.assertEqual(r['consolidated'], [])

    def test_conductor_numeric(self):
        c = self.config(); c['rules'] = [self.rule()]
        q = run(self.index, c)['quantities'][1]
        self.assertAlmostEqual(q['quantity'], 26.4)
        self.assertEqual(q['method'], 'calculated'); self.assertEqual(q['basis_id'], 'route')

    def test_rule_approval(self):
        c = self.config(); c['rules'] = [self.rule(status='proposed')]
        self.assertIsNone(run(self.index, c)['quantities'][1]['quantity'])

    def test_rule_missing_variable(self):
        c = self.config(); rule = self.rule(); del rule['variables']['vertical']; c['rules'] = [rule]
        self.assertIsNone(run(self.index, c)['quantities'][1]['quantity'])

    def test_rule_dimensional_units(self):
        c = self.config(); rule = self.rule(); rule['variable_units']['vertical'] = 'mm'; c['rules'] = [rule]
        self.assertIsNone(run(self.index, c)['quantities'][1]['quantity'])

    def test_rule_condition(self):
        c = self.config(); c['rules'] = [self.rule(conditions={'system': 'other'})]
        self.assertIsNone(run(self.index, c)['quantities'][1]['quantity'])

    def test_no_derived_from_disputed_base(self):
        c = self.config(); c['selections'].append(self.selection(id='second')); c['rules'] = [self.rule()]
        self.assertIsNone(run(self.index, c)['quantities'][-1]['quantity'])

    def test_support_rule(self):
        c = self.config(); c['rules'] = [self.rule(kind='supports', unit='u',
            variables={'spacing': 2, 'endpoints': 1}, variable_units={'spacing': 'm', 'endpoints': 'u'}, single_run_verified=True)]
        self.assertEqual(run(self.index, c)['quantities'][1]['quantity'], 4)

    def test_cross_document_contradiction(self):
        idx = copy.deepcopy(self.index); idx['cells'] = [{'id': 'cell', 'value': 6, 'document_id': 'reference', 'sheet': 'Qty', 'address': 'C10'}]
        c = self.config(); c['rules'] = [self.rule()]; c['comparisons'] = [{'quantity_id': 'route', 'cell_id': 'cell', 'unit': 'm', 'absolute_tolerance': 0, 'status': 'approved', 'evidence': 'Same synthetic scope'}]
        r = run(idx, c); self.assertEqual(r['consolidated'], [])
        self.assertTrue(any(o['code'] == 'DOCUMENT_CONTRADICTION' for o in r['observations']))

    def test_same_name_not_duplicate(self):
        c = self.config(); c['selections'] = [self.selection(kind='count', unit='u', source_refs=[self.words[i]['id']], tag_is_element_verified=True, id=f'lamp{i}') for i in range(2)]
        r = run(self.index, c); self.assertEqual(r['consolidated'][0]['quantity'], 2)

    def test_revision_family_conflict(self):
        idx = copy.deepcopy(self.index); d = copy.deepcopy(self.doc); d['id'] = 'other'; d['sha256'] = 'other'; idx['documents'].append(d)
        c = self.config(); setting = copy.deepcopy(c['documents'][self.doc['id']]); setting.update(sha256='other', revision='R02'); c['documents']['other'] = setting
        self.assertEqual(run(idx, c)['consolidated'], [])

    def test_rotated_pdf_coordinate_frame(self):
        p = self.index['pages'][1]
        self.assertEqual(p['original_rotation'], 90); self.assertEqual(p['width'], 300)
        ref = next(r for r in self.index['references'] if r['page_id'] == p['id'])
        self.assertEqual(geometry_length(ref['commands']), 50)

    def test_traceability_complete(self):
        c = self.config(); c['rules'] = [self.rule()]; r = run(self.index, c)
        for q in r['quantities']:
            for key in ['id', 'project', 'document_id', 'page_id', 'revision', 'source_refs', 'formula', 'evidence', 'confidence', 'review_status']:
                self.assertTrue(q[key], key)

    def test_no_mutation_originals(self):
        before = (self.inputs / self.doc['name']).read_bytes()
        run(self.index, self.config())
        self.assertEqual((self.inputs / self.doc['name']).read_bytes(), before)

    def test_outputs_and_formula_graph(self):
        from openpyxl import load_workbook
        c = self.config(); c['rules'] = [self.rule()]; r = run(self.index, c)
        out = self.root / 'results'; export_all(r, out, index_directory=self.index_dir)
        w = load_workbook(out / 'takeoff.xlsx', data_only=False)
        self.assertEqual(w.sheetnames, SHEETS)
        self.assertIn("'Trazabilidad'!K2", w['Materiales derivados']['D2'].value)
        self.assertEqual(w['Trazabilidad']['K3'].value, "='Materiales derivados'!D2")
        self.assertEqual(w['Computo consolidado'].max_row, 3)
        w.close(); self.assertTrue((out / self.page['image']).exists())
        self.assertIn('const data=JSON.parse', (out / 'viewer.html').read_text())

    def test_output_cannot_write_into_originals(self):
        with self.assertRaises(ValueError): index_folder(self.inputs, self.inputs / 'output', 'bad')

    def test_malicious_text_escaped(self):
        from openpyxl import load_workbook
        r = run(self.index, self.config()); r['project'] = '</script><script>alert(1)</script>'
        r['quantities'][0]['description'] = '=HYPERLINK("http://evil")'
        out = self.root / 'malicious'; export_all(r, out, index_directory=self.index_dir)
        html = (out / 'viewer.html').read_text()
        self.assertNotIn('</script><script>alert(1)', html)
        w = load_workbook(out / 'takeoff.xlsx'); self.assertEqual(w['Trazabilidad']['I2'].data_type, 's'); w.close()

    def test_dwg_blocked_not_read_as_dxf(self):
        folder = self.root / 'dwg'; folder.mkdir(exist_ok=True)
        (folder / 'SYNTHETIC.dwg').write_bytes(b'AC1032' + b'\0' * 20)
        r = index_folder(folder, self.root / 'dwg-out', 'SYNTHETIC')
        self.assertEqual(r['documents'][0]['status'], 'pending_adapter')
        self.assertEqual(r['documents'][0]['dwg_signature'], 'AC1032'); self.assertEqual(r['references'], [])

    def test_cli_hash_change_rejected(self):
        folder = self.root / 'changed'; folder.mkdir(exist_ok=True)
        (folder / 'note.txt').write_text('before'); idx = index_folder(folder, self.root / 'changed-index', 'SYNTHETIC')
        save(self.root / 'changed-index' / 'index.json', idx); save(self.root / 'empty-config.json', {'schema': 'eqt/1'})
        (folder / 'note.txt').write_text('after')
        r = subprocess.run([sys.executable, str(Path(__file__).resolve().parents[1] / 'takeoff.py'), 'compute',
                            str(self.root / 'changed-index/index.json'), '--config', str(self.root / 'empty-config.json'),
                            '--out', str(self.root / 'changed-result')], capture_output=True, text=True)
        self.assertEqual(r.returncode, 2); self.assertIn('cambió', r.stderr)

    def test_duplicate_calibration_invalidates_all(self):
        c = self.config(); c['calibrations'].append(copy.deepcopy(c['calibrations'][0]))
        self.assertIsNone(run(self.index, c)['quantities'][0]['quantity'])

    def test_geometric_duplicate_inside_one_route(self):
        idx = copy.deepcopy(self.index); ref = copy.deepcopy(self.paths[0]); ref['id'] = 'copy'
        idx['references'].append(ref); c = self.config(); c['selections'][0]['source_refs'].append('copy')
        self.assertIsNone(run(idx, c)['quantities'][0]['quantity'])

    def test_coverage_stays_partial(self):
        r = run(self.index, self.config())
        self.assertEqual(next(c for c in r['coverage'] if c['system'] == 'bandejas')['status'], 'partial')
        self.assertEqual(next(c for c in r['coverage'] if c['system'] == 'respaldo')['status'], 'not_reviewed')

    def test_pending_unquantified_material_not_omitted(self):
        c = self.config(); c['selections'].append(self.selection(id='boxes', kind='pending', unit='u',
            system='iluminacion', description='Cajas necesarias', material_key='box', source_refs=[], reason='Falta detalle de montaje'))
        q = run(self.index, c)['quantities'][-1]
        self.assertIsNone(q['quantity']); self.assertIn('Falta detalle', q['notes'][0])

    def test_manual_comparison_metrics(self):
        from benchmark import compare
        p = {('a', 'bandejas', 'tray', 'm'): 11, ('extra', 'bandejas', 'tray', 'm'): 1}
        t = {('a', 'bandejas', 'tray', 'm'): 10, ('missing', 'bandejas', 'tray', 'm'): 2}
        r = compare(p, t)['bandejas']
        self.assertEqual(r['false_positive_identities'], 1); self.assertEqual(r['false_negative_identities'], 1)
        self.assertAlmostEqual(r['length_errors'][0]['relative_error'], .1)

    def test_pending_keeps_graphic_provenance(self):
        c = self.config(); c['selections'][0].update(kind='pending', reason='Detalle incompleto')
        q = run(self.index, c)['quantities'][0]
        self.assertEqual(q['page_id'], self.page['id']); self.assertEqual(q['document_id'], self.doc['id'])
        self.assertIsNone(q['quantity'])

    def test_third_party_responsibility_is_preserved(self):
        c = self.config(); c['selections'][0].update(provision_by='cliente', installation_by='contratista')
        q = run(self.index, c)['quantities'][0]
        self.assertEqual(q['provision_by'], 'cliente'); self.assertEqual(q['installation_by'], 'contratista')
        self.assertEqual(q['quantity'], 5)

    def test_scale_error_explains_missing_calibration(self):
        c = self.config(); c['calibrations'] = []
        q = run(self.index, c)['quantities'][0]
        self.assertIn('Falta calibración válida', q['notes'][0])


if __name__ == '__main__':
    unittest.main()
